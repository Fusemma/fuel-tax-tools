"""Builds the state info workbook.

One list per kind of info (states, returns, licenses, rates, contacts,
tax setups), every row tied to a state by its two-letter code, and a
State page that shows one state from all the lists at once.

Georgia is filled in as the worked example, public facts only, each with
the source it came from. Anything about the company (licenses held,
numbers, bond amounts, where files live, tax setups) gets typed in at
work.

Build, then recalc so the saved file carries values:
  python build_workbook.py
  powershell -NoProfile -File "..\\GA Sales\\recalc.ps1" "<full path to State Info.xlsx>"
"""
import datetime as dt
import os

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill, Protection
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.pagebreak import Break
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "State Info.xlsx")
LAST = 1000  # the State page reads each list down to this row
CHECKED = dt.date(2026, 10, 10)

F = lambda **kw: Font(name="Arial", size=10, **kw)
BOLD = F(bold=True)
TITLE = Font(name="Arial", size=16, bold=True)
SECTION = Font(name="Arial", size=11, bold=True, color="FFFFFF")
SECTION_FILL = PatternFill("solid", fgColor="44546A")
LABEL_FILL = PatternFill("solid", fgColor="DDDDDD")
YELLOW = PatternFill("solid", fgColor="FFFF00")
PICK = Font(name="Arial", size=14, bold=True, color="0000FF")
LINK = F(color="0563C1", underline="single")
NOTE = F(color="808080", italic=True)
TOP = Alignment(vertical="top")
WRAP = Alignment(wrap_text=True, vertical="top")
DATE = "m/d/yyyy"

STATES = [
    ("AL", "Alabama"), ("AK", "Alaska"), ("AZ", "Arizona"), ("AR", "Arkansas"),
    ("CA", "California"), ("CO", "Colorado"), ("CT", "Connecticut"),
    ("DE", "Delaware"), ("DC", "District of Columbia"), ("FL", "Florida"),
    ("GA", "Georgia"), ("HI", "Hawaii"), ("ID", "Idaho"), ("IL", "Illinois"),
    ("IN", "Indiana"), ("IA", "Iowa"), ("KS", "Kansas"), ("KY", "Kentucky"),
    ("LA", "Louisiana"), ("ME", "Maine"), ("MD", "Maryland"),
    ("MA", "Massachusetts"), ("MI", "Michigan"), ("MN", "Minnesota"),
    ("MS", "Mississippi"), ("MO", "Missouri"), ("MT", "Montana"),
    ("NE", "Nebraska"), ("NV", "Nevada"), ("NH", "New Hampshire"),
    ("NJ", "New Jersey"), ("NM", "New Mexico"), ("NY", "New York"),
    ("NC", "North Carolina"), ("ND", "North Dakota"), ("OH", "Ohio"),
    ("OK", "Oklahoma"), ("OR", "Oregon"), ("PA", "Pennsylvania"),
    ("RI", "Rhode Island"), ("SC", "South Carolina"), ("SD", "South Dakota"),
    ("TN", "Tennessee"), ("TX", "Texas"), ("UT", "Utah"), ("VT", "Vermont"),
    ("VA", "Virginia"), ("WA", "Washington"), ("WV", "West Virginia"),
    ("WI", "Wisconsin"), ("WY", "Wyoming"),
]

# Georgia sources, all official
GA_FAQ = "https://dor.georgia.gov/2026-suspension-georgia-motor-fuel-taxes-faqs"
GA_DYED = "https://dor.georgia.gov/mft-2026-04-temporary-waiver-state-dyed-fuel-penalties"
GA_PREPAID = "https://dor.georgia.gov/prepaid-local-tax-motor-fuel-sales-effective-july-1-2026"
GA_CONTACT = "https://dor.georgia.gov/find-contact-information-topic"
GA_ST3 = "DOR ST-3 instructions, rev 07/25/22"
GA_MF12 = "DOR form MF-12, rev 06/01/17"

# Each list: its columns (name, width) and its rows. State always comes first.
LISTS = {
    "States": (
        [("State", 7), ("Name", 18), ("Tax agency", 30), ("Website", 28),
         ("Filing site", 26), ("Main phone", 16), ("Bond required", 14),
         ("Where the files live", 32), ("Quirks", 60), ("Source", 44),
         ("Last checked", 13)],
        [["GA", "Georgia", "Georgia Department of Revenue",
          "https://dor.georgia.gov", "https://gtc.dor.ga.gov", "(877) 423-6711",
          None, None,
          "Two sales tax returns, on-road and off-road, both laid out like the "
          "ST-3. On-road is gallons by county, two lines per county, and Fulton "
          "County splits by rate. State excise suspended Sept 29 to Nov 5, 2026, "
          "and state penalties for dyed diesel on the road waived Oct 6 to Nov 5.",
          GA_CONTACT, CHECKED]
         if code == "GA" else [code, name] + [None] * 9
         for code, name in STATES]),
    "Returns": (
        [("State", 7), ("Return", 30), ("Form", 22), ("Tax", 28),
         ("How often", 12), ("How to file", 26), ("How to pay", 18),
         ("Notes", 60), ("Source", 44), ("Last checked", 13)],
        [["GA", "Sales and use tax, on-road", "ST-3 layout",
          "Sales tax, mostly prepaid local", "Monthly", "Georgia Tax Center", None,
          "Gallons by county and product, two lines per county. College Park "
          "goes on line 4, every other county on line 6. Fulton County splits "
          "by rate.", GA_ST3, CHECKED],
         ["GA", "Sales and use tax, off-road", "ST-3 layout", "Sales and use tax",
          "Monthly", "Georgia Tax Center", None,
          "Dyed fuel. Georgia taxes the small per-gallon fees too, so taxable "
          "sales are product cost plus fees.", GA_ST3, CHECKED],
         ["GA", "Motor fuel excise, distributor return", None,
          "State motor fuel excise", "Monthly", "Georgia Tax Center", None,
          "Schedules of Receipts and Disbursements. Sept 29 to Nov 5, 2026: "
          "report the gallons on both schedules, don't remit the excise.",
          GA_FAQ, CHECKED],
         ["GA", "Refund claim", "MFD-33, with MF-12 waivers",
          "Motor fuel excise refunds", "As needed",
          "Georgia Tax Center, as refund attachments", None,
          "One MF-12 per supplier per refund type, signed by the supplier "
          "before a notary. Each load's invoice copies go in merged into one PDF.",
          GA_MF12, CHECKED]]),
    "Licenses": (
        [("State", 7), ("License", 30), ("Number", 18), ("Expires", 12),
         ("Bond amount", 14), ("Renew how", 24), ("Notes", 50), ("Source", 30),
         ("Last checked", 13)],
        [["GA", "Motor fuel distributor", None, None, None, None,
          "Type the number and bond in at work.", None, None]]),
    "Rates": (
        [("State", 7), ("Product", 20), ("Tax", 26), ("Rate", 10), ("Per", 20),
         ("From", 12), ("To", 12), ("Notes", 56), ("Source", 44),
         ("Last checked", 13)],
        [["GA", "Gasoline", "State excise", 0.333, "$ a gallon", None, None,
          "The rate at the Sept 2026 suspension. Georgia resets it every Jan 1.",
          GA_FAQ, CHECKED],
         ["GA", "Diesel", "State excise", 0.373, "$ a gallon", None, None,
          "The rate at the Sept 2026 suspension. Georgia resets it every Jan 1.",
          GA_FAQ, CHECKED],
         ["GA", "Gasoline and diesel", "State excise, suspended", 0, "$ a gallon",
          dt.date(2026, 9, 29), dt.date(2026, 11, 5),
          "Governor's order EO 09.28.26.01, extended to Nov 5 by EO 10.06.26.01. "
          "Gallons still get reported. Prepaid local tax isn't suspended.",
          GA_FAQ, CHECKED],
         ["GA", "Gasoline and diesel", "Prepaid local sales tax", None,
          "see notes", None, None,
          "The local rate times a state-set price per gallon, reset each Jan 1 "
          "and July 1 and never over $3.00. $3.00 for gas and diesel from "
          "July 1, 2026.", GA_PREPAID, CHECKED],
         ["GA", "Dyed diesel, road use", "State penalties waived",
          None, "see notes", dt.date(2026, 10, 6), dt.date(2026, 11, 5),
          "EO 10.06.26.01 and Bulletin MFT-2026-04 waive the 48-9-16(e) "
          "penalties for dyed fuel on the road. The $100 a month notice penalty "
          "in (f) still applies.",
          GA_DYED, CHECKED]]),
    "Contacts": (
        [("State", 7), ("Who", 30), ("Phone", 16), ("Email", 28), ("Hours", 18),
         ("What for", 40), ("Source", 44), ("Last checked", 13)],
        [["GA", "Taxpayer Services, motor fuel", "(877) 423-6711", None,
          "8 to 5, weekdays", "Motor fuel tax questions", GA_CONTACT, CHECKED]]),
    "Setups": (
        [("State", 7), ("Setup", 28), ("Product", 20), ("What it charges", 30),
         ("Code in the system", 20), ("Notes", 50), ("Last checked", 13)],
        []),
}

# What the State page shows from each list: (field, page column), and how
# many rows it has room for.
SECTIONS = [
    ("Returns", "Returns", [("Return", "B"), ("Form", "C"), ("Tax", "D"),
                            ("How often", "E"), ("How to file", "F"),
                            ("How to pay", "G"), ("Notes", "H")], 8),
    ("Licenses", "Licenses", [("License", "B"), ("Number", "C"), ("Expires", "D"),
                              ("Bond amount", "E"), ("Renew how", "F"),
                              ("Notes", "H")], 6),
    # only rates with no To date, or a To date still ahead; old ones stay on the tab
    ("Rates", "Current rates", [("Product", "B"), ("Tax", "C"), ("Rate", "D"),
                                ("Per", "E"), ("From", "F"), ("To", "G"),
                                ("Notes", "H")], 8),
    ("Contacts", "Contacts", [("Who", "B"), ("Phone", "C"), ("Email", "D"),
                              ("Hours", "E"), ("What for", "F")], 5),
    ("Setups", "Tax setups", [("Setup", "B"), ("Product", "C"),
                              ("What it charges", "D"), ("Code in the system", "E"),
                              ("Notes", "H")], 6),
]
SLOT_HEIGHT = 27  # two lines; Excel won't grow a row when a formula's text changes
FORMATS = {"Expires": DATE, "From": DATE, "To": DATE, "Last checked": DATE,
           "Bond amount": "$#,##0"}
LONG_TEXT = {"Quirks", "Notes", "What for", "Return", "Tax", "How to file", "Form", "Per"}

wb = openpyxl.Workbook()
wb.remove(wb.active)
page = wb.create_sheet("State")
page.sheet_properties.tabColor = "70AD47"


def col_of(list_name, field):
    names = [name for name, _ in LISTS[list_name][0]]
    return "ABCDEFGHIJKLMNOP"[names.index(field)]


def link_or_text(value):
    # a web address or a folder path becomes a link, anything else stays text
    return (f'IF(OR(LEFT({value},4)="http",LEFT({value},2)="\\\\",MID({value},2,1)=":"),'
            f'HYPERLINK({value},{value}),{value})')


# --- the lists ---
for name, (columns, rows) in LISTS.items():
    ws = wb.create_sheet(name)
    for i, (heading, width) in enumerate(columns, 1):
        cell = ws.cell(row=1, column=i, value=heading)
        cell.font = BOLD
        ws.column_dimensions["ABCDEFGHIJKLMNOP"[i - 1]].width = width
    for r, row in enumerate(rows or [[None] * len(columns)], 2):
        for i, value in enumerate(row, 1):
            heading = columns[i - 1][0]
            cell = ws.cell(row=r, column=i, value=value)
            cell.font = F()
            cell.alignment = WRAP if heading in LONG_TEXT else TOP
            if heading in FORMATS:
                cell.number_format = FORMATS[heading]
            if isinstance(value, str) and value.startswith("http"):
                cell.hyperlink = value
                cell.font = LINK
    end = "ABCDEFGHIJKLMNOP"[len(columns) - 1]
    table = Table(displayName=name + "List", ref=f"A1:{end}{max(len(rows), 1) + 1}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(table)
    ws.freeze_panes = "B2"
    if name != "States":
        pick = DataValidation(type="list", formula1="StateList", allow_blank=True,
                              showErrorMessage=True, errorTitle="State code",
                              error="Pick a two-letter state code from the list.")
        ws.add_data_validation(pick)
        pick.add(f"A2:A{LAST}")

wb.defined_names["StateList"] = DefinedName("StateList", attr_text=f"States!$A$2:$A${len(STATES) + 1}")

# --- the State page ---
for letter, width in zip("ABCDEFGH", [4, 30, 20, 24, 14, 22, 14, 70]):
    page.column_dimensions[letter].width = width
page.column_dimensions["A"].hidden = True  # helper numbers live here

page["B1"] = "State info"
page["B1"].font = TITLE
page["B2"] = ("Pick a state in the yellow box. Everything below comes from the "
              "other tabs; make changes there.")
page["B2"].font = NOTE
page["B3"] = "State"
page["B3"].font = BOLD
page["B3"].alignment = Alignment(vertical="center")
page["C3"] = "GA"
page["C3"].font = PICK
page["C3"].fill = YELLOW
page["C3"].alignment = Alignment(horizontal="center", vertical="center")
page["C3"].protection = Protection(locked=False)
picker = DataValidation(type="list", formula1="StateList", allow_blank=False,
                        showErrorMessage=True, errorTitle="State code",
                        error="Pick a two-letter state code from the list.")
page.add_data_validation(picker)
picker.add("C3")
page.row_dimensions[3].height = 24

# A4: which States row the picked state is on
page["A4"] = f"=IFERROR(MATCH($C$3,States!$A$2:$A${LAST},0),\"\")"
page["D3"] = f'=IF($A$4="","",INDEX(States!$B$2:$B${LAST},$A$4))'
page["D3"].font = TITLE
page["D3"].alignment = Alignment(vertical="center")


def section_header(row, text):
    for letter in "BCDEFGH":
        page[f"{letter}{row}"].fill = SECTION_FILL
    page[f"B{row}"] = text
    page[f"B{row}"].font = SECTION


row = 5
section_header(row, "Basics")
for field in ["Tax agency", "Website", "Filing site", "Main phone", "Bond required",
              "Where the files live", "Quirks", "Source", "Last checked"]:
    row += 1
    rng = f"States!{col_of('States', field)}$2:{col_of('States', field)}${LAST}"
    value = f"INDEX({rng},$A$4)"
    shown = link_or_text(value) if field in {"Website", "Filing site", "Where the files live", "Source"} else value
    page[f"B{row}"] = field
    page[f"B{row}"].font = BOLD
    page[f"B{row}"].fill = LABEL_FILL
    page[f"B{row}"].alignment = TOP
    page.merge_cells(f"C{row}:H{row}")
    cell = page[f"C{row}"]
    cell.value = f'=IF($A$4="","",IF({value}="","",{shown}))'
    cell.font = LINK if field in {"Website", "Filing site", "Where the files live", "Source"} else F()
    cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
    if field in FORMATS:
        cell.number_format = FORMATS[field]
    if field == "Quirks":
        page.row_dimensions[row].height = SLOT_HEIGHT

for list_name, title, fields, slots in SECTIONS:
    row += 2
    states_col = f"{list_name}!$A$2:$A${LAST}"
    match = f"({states_col}=$C$3)"
    count = f"COUNTIF({states_col},$C$3)"
    if list_name == "Rates":
        to = f"Rates!${col_of('Rates', 'To')}$2:${col_of('Rates', 'To')}${LAST}"
        match = f'(({states_col}=$C$3)*(({to}="")+({to}>=TODAY())))'
        count = f"SUMPRODUCT({match})"
        page.row_breaks.append(Break(id=row - 1))  # rates start the second printed page
    section_header(row, None)
    page[f"B{row}"] = f'="{title} ("&{count}&")"'
    page[f"B{row}"].font = SECTION
    row += 1
    for field, letter in fields:
        page[f"{letter}{row}"] = field
        page[f"{letter}{row}"].font = BOLD
    for letter in "BCDEFGH":
        page[f"{letter}{row}"].fill = LABEL_FILL
    for k in range(1, slots + 1):
        row += 1
        # A: the list row of this state's k-th entry, blank when there isn't one
        page[f"A{row}"] = (f"=IFERROR(_xlfn.AGGREGATE(15,6,(ROW({states_col})-1)"
                           f"/{match},{k}),\"\")")
        page.row_dimensions[row].height = 15 if list_name == "Contacts" else SLOT_HEIGHT
        for field, letter in fields:
            col = col_of(list_name, field)
            value = f"INDEX({list_name}!{col}$2:{col}${LAST},$A{row})"
            cell = page[f"{letter}{row}"]
            cell.value = f'=IF($A{row}="","",IF({value}="","",{value}))'
            cell.font = F()
            cell.alignment = WRAP if field in LONG_TEXT else TOP
            if field in FORMATS:
                cell.number_format = FORMATS[field]
    row += 1
    page[f"B{row}"] = (f'=IF({count}>{slots},"Plus "&({count}-{slots})&'
                       f'" more on the {list_name} tab","")')
    page[f"B{row}"].font = NOTE

page.freeze_panes = "A5"
page.print_area = f"B1:H{row}"
page.page_setup.orientation = "landscape"
page.page_setup.fitToWidth = 1
page.page_setup.fitToHeight = 0
page.sheet_properties.pageSetUpPr.fitToPage = True
page.print_title_rows = "1:3"
page.protection.sheet = True  # no password; Review, Unprotect Sheet to change it
page.protection.formatColumns = False
page.protection.formatRows = False

# --- how to use ---
how = wb.create_sheet("How to use")
how.column_dimensions["B"].width = 100
lines = [
    ("How this workbook works", TITLE),
    ("The State tab shows one state at a time. Pick the state in its yellow box "
     "and everything about that state shows up, pulled from the other tabs.", None),
    ("Only the yellow box on the State tab takes typing. Type everything else on "
     "the tab it belongs to: States, Returns, Licenses, Rates, Contacts or Setups.", None),
    ("Every row starts with the state's two-letter code. The State column only "
     "takes codes from the list, so spellings can't drift apart.", None),
    ("A state with more of something just gets more rows. Four Georgia returns "
     "are four rows on Returns.", None),
    ("To add a row, type in the first empty row right under a list, and the list "
     "grows to take it.", None),
    ("A new kind of info is a new column at the right end of that tab's list. "
     "It shows on the State tab once someone adds it there too.", None),
    ("Rates keep their history. When a rate changes, give the old row a To date "
     "and add a new row with its From date. The State tab shows only current "
     "rates: no To date, or a To date still ahead. Old ones stay on the Rates tab.", None),
    ("Source and Last checked: put the page you checked and the date, so anyone "
     "can recheck it in seconds.", None),
    ("The State tab has room for 8 returns, 6 licenses, 8 current rates, 5 "
     "contacts and 6 tax setups per state, and says so when a state has more. "
     "It shows two lines of each note; the whole note is on its tab.", None),
    ("The State tab is locked so nobody types over it by accident. No password: "
     "Review, then Unprotect Sheet, if it ever needs changing.", None),
    ("Company numbers, like license and account numbers and bond amounts, get "
     "typed in here at work.", None),
    ("Never put passwords or logins in this workbook.", None),
]
for r, (text, font) in enumerate(lines, 1):
    cell = how.cell(row=r * 2 - 1, column=2, value=text)
    cell.font = font or F()
    cell.alignment = Alignment(wrap_text=True, vertical="top")

wb.active = 0
wb.properties.creator = ""
wb.save(OUT)
print("built", OUT)
