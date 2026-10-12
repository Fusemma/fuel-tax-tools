# Builds the federal excise exemption certificate template: IRS Model
# Certificates M and P joined into one form, the way most suppliers send
# it. A state or local government (or a nonprofit school, for gasoline)
# signs it so the vendor can claim back the federal excise tax.
# Rebuild any time with:  python build_exemption_template.py
# Output lands beside this script. The wording is set out line by line in
# "Combined certificate text (Pub 510 Dec 2025).txt" here, from Pub 510's
# Model M (page 53) and Model P (page 56); the layout follows Model M's
# page. Same shape as the 637 template: every fill-in spot is a
# click-and-type field showing gray hint text until something is typed in
# it, the boxes check with a click, and the signature line stays blank
# for ink.
import datetime
import os

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Inches, Pt, RGBColor, Twips

from word_fields import checkbox, field, finish, line_under

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, 'Exemption Certificate - TEMPLATE.docx')

BOX_WIDTH = Inches(7.3)   # the border around the certificate, margin to margin
GUTTER = Inches(0.4)      # strip left of the text where the boxes sit
GAP = Pt(4)               # space under each line of the certificate
FUELS = 'gasoline, aviation gasoline, diesel fuel, or kerosene'

doc = Document()

section = doc.sections[0]
section.page_width, section.page_height = Inches(8.5), Inches(11)
section.left_margin = section.right_margin = Inches(0.6)
section.top_margin = section.bottom_margin = Inches(0.5)

# Open as a current Word file (no "Compatibility Mode"), at 100% zoom
settings = doc.settings.element
for setting in settings.iter(qn('w:compatSetting')):
    if setting.get(qn('w:name')) == 'compatibilityMode':
        setting.set(qn('w:val'), '15')
zoom = settings.find(qn('w:zoom'))
zoom.attrib.clear()
zoom.set(qn('w:percent'), '100')

# No author or program name in the file's properties
props = doc.core_properties
props.author = props.last_modified_by = props.comments = ''
props.created = props.modified = datetime.datetime.now(
    datetime.timezone.utc).replace(tzinfo=None, microsecond=0)
props.revision = 1

normal = doc.styles['Normal']
normal.font.name = 'Arial'
normal.font.size = Pt(10)
normal.paragraph_format.space_before = Pt(0)
normal.paragraph_format.space_after = Pt(0)
normal.paragraph_format.line_spacing = 1

ph = doc.styles.add_style('Placeholder Text', WD_STYLE_TYPE.CHARACTER)
ph.font.color.rgb = RGBColor(0x80, 0x80, 0x80)

# The certificate sits in one bordered box
box = doc.add_table(rows=1, cols=1)
box._tbl.replace(box._tbl.tblPr, parse_xml(
    '<w:tblPr {ns}>'
    '<w:tblW w:w="{w}" w:type="dxa"/>'
    '<w:tblInd w:w="0" w:type="dxa"/>'
    '<w:tblBorders>'
    '<w:top w:val="single" w:sz="6" w:space="0" w:color="000000"/>'
    '<w:left w:val="single" w:sz="6" w:space="0" w:color="000000"/>'
    '<w:bottom w:val="single" w:sz="6" w:space="0" w:color="000000"/>'
    '<w:right w:val="single" w:sz="6" w:space="0" w:color="000000"/>'
    '</w:tblBorders>'
    '<w:tblLayout w:type="fixed"/>'
    '<w:tblCellMar>'
    '<w:top w:w="72" w:type="dxa"/><w:left w:w="173" w:type="dxa"/>'
    '<w:bottom w:w="115" w:type="dxa"/><w:right w:w="173" w:type="dxa"/>'
    '</w:tblCellMar>'
    '</w:tblPr>'.format(ns=nsdecls('w'), w=BOX_WIDTH.twips)))
box.columns[0].width = BOX_WIDTH
cell = box.cell(0, 0)
cell.width = BOX_WIDTH

p = cell.paragraphs[0]
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run('CERTIFICATE FOR STATE USE OR NONPROFIT EDUCATIONAL')
run.bold = True
run.add_break()
p.add_run('ORGANIZATION USE').bold = True
p.paragraph_format.space_after = Pt(4)

p = cell.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run('(To support vendor’s claim for a credit or payment')
run.bold = True
run.add_break()
p.add_run('under § 6416(a)(4) and/or § 6427 of the Internal Revenue '
          'Code.)').bold = True

# Vendor block: three lines to write on over one label. Each line's right
# edge sits a twip apart from the next; without that, Word joins
# same-bordered paragraphs and draws only the last line.
for n, hint in enumerate(('Vendor name', 'Vendor address', 'Vendor FEIN')):
    p = cell.add_paragraph()
    line_under(p)
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.right_indent = Twips(n)
    field(p, hint)
p = cell.add_paragraph(
    'Name, address, and employer identification number of ultimate vendor')
p.paragraph_format.space_before = Pt(2)
p.paragraph_format.space_after = Pt(6)


def body(text='', box_title=None):
    """A certificate line, indented; a box in the strip to its left if asked."""
    p = cell.add_paragraph()
    p.paragraph_format.left_indent = GUTTER
    p.paragraph_format.space_after = GAP
    if box_title:
        p.paragraph_format.first_line_indent = -GUTTER
        checkbox(p, box_title)
        p.add_run('\t')
    if text:
        p.add_run(text)
    return p


body('The undersigned ultimate purchaser (“Buyer”) hereby certifies the '
     'following under the penalties of perjury:')
body('Buyer will use the ' + FUELS + ' to which this certificate relates '
     '(check one):')   # no-break space keeps "check one" on one line
body('For the exclusive use of a state or local government, or the District '
     'of Columbia; or', 'State or local government')
body('For the exclusive use of a nonprofit educational organization '
     '(gasoline or aviation gasoline only).',
     'Nonprofit educational organization')
body('This certificate applies to the following (complete as applicable):')

body('This is a single purchase certificate:', 'Single purchase')
p = body('1. ')
field(p, 'Invoice no.', 'Invoice or delivery ticket number')
p.add_run(' Invoice or delivery ticket number')
p = body('2. ')
field(p, 'Gallons', 'Number of gallons')
p.add_run(' Number of gallons')

body('This is a certificate covering all purchases under a specified '
     'account or order number:', 'All purchases on an account')
p = body('1. Effective date ')
field(p, 'MM/DD/YYYY', 'Effective date')
p = body('2. Expiration date ')
field(p, 'MM/DD/YYYY', 'Expiration date')
p.add_run(' (period not to exceed 1 year after the effective date)')
p = body('3. Buyer’s account or order number ')
field(p, 'Account or order no.', 'Buyer’s account or order number')

body('Buyer will provide a new certificate to the vendor if any information '
     'in this certificate changes.')
body('Buyer understands that by signing this certificate, Buyer gives up its '
     'right to claim any credit or payment for the ' + FUELS + ' to which '
     'this certificate relates.')
body('If Buyer uses the diesel fuel or kerosene to which this certificate '
     'relates for a purpose other than stated in the certificate, Buyer will '
     'be liable for any tax.')
body('Buyer acknowledges that it hasn’t and won’t claim any credit or '
     'payment for the ' + FUELS + ' to which this certificate relates.')
body('Buyer understands that the fraudulent use of this certificate may '
     'subject Buyer and all parties making such fraudulent use of this '
     'certificate to a fine or imprisonment, or both, together with the '
     'costs of prosecution.')

# Signature block: each fill-in (or ink) line over its label. The
# signature line gets extra room above it for a pen.
rows = [
    ('Name of person signing', 'Printed or typed name of person signing'),
    ('Title of person signing', 'Title of person signing'),
    ('Buyer name', 'Name of Buyer'),
    ('Buyer FEIN', 'Employer identification number'),
    ('Buyer address', 'Address of Buyer'),
    (None, 'Signature and date signed'),
]
for hint, label in rows:
    p = cell.add_paragraph()
    line_under(p)
    p.paragraph_format.space_before = Pt(18 if hint is None else 7)
    if hint:
        field(p, hint)
    p = cell.add_paragraph(label)
    p.paragraph_format.space_before = Pt(1)

# Word needs a paragraph after a table; keep it tiny so it can't push a
# second page.
p = doc.add_paragraph()
p.paragraph_format.line_spacing = Pt(1)

finish(doc)
doc.save(OUT)

# Read it back as a sanity check
check = Document(OUT)
sdts = check.element.body.findall('.//' + qn('w:sdt'))
boxes = [s for s in sdts if s.find(
    './/{http://schemas.microsoft.com/office/word/2010/wordml}checkbox')
    is not None]
print('saved:', OUT)
print('fields:', len(sdts) - len(boxes))
print('boxes:', len(boxes))
